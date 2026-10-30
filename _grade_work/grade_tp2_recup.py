# -*- coding: utf-8 -*-
"""Corrige les TP2 récupérés et écrit un Excel mis à jour SANS écraser l'original."""
from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

import build_notes_excel as b
from openpyxl import load_workbook
from openpyxl.styles import Font

TP_RECUP = Path(r"c:\Users\mhado\Desktop\TP récupéré")
STAGING = Path(__file__).resolve().parent / "tp2_recup_staging"
ORIG_XLSX = b.OUT_XLSX
NEW_XLSX = ORIG_XLSX.with_name("Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx")

# (clé_soumission, liste de fichiers/dossiers relatifs à TP_RECUP, élèves associés NOM_Prenom keys)
ASSIGNMENTS = [
    {
        "key": "BOUCHER_JULES",
        "also": [],
        "paths": ["Boucher/Tp2 (1).docx"],
        "label": "Boucher/Tp2 (1).docx",
    },
    {
        "key": "NORMAND_MAXIME",
        "also": ["FRANCOIS_ALBAN"],
        "paths": ["Normand_Francois"],
        "label": "Normand_Francois/ (TP2.docx + IMG_1620.mov)",
    },
    {
        "key": "BENOIT_CONSTANCE",
        "also": ["ANCENAY_ROMAIN"],
        "paths": ["SIN_TP2_Constance_BENOIT_Romain_ANCENAY.zip"],
        "label": "SIN_TP2_Constance_BENOIT_Romain_ANCENAY.zip",
    },
    {
        "key": "HOEFFLER_LEO",
        "also": ["WIPFF_HUGO"],
        "paths": None,  # resolved by glob
        "glob": "TP_2_HOEFFLER*.zip",
        "label": "TP_2_HOEFFLER_Léo_WIPFF_Hugo_1STI2D1.zip",
    },
    {
        "key": "COCOL_ALEXIS",
        "also": ["DIOWO_CORENTIN"],
        "paths": None,
        "glob": "TP_2_Mesurer*.docx",
        "label": "TP_2_Mesurer_la_luminosité_avec_une_LDR.docx (Alexis & Corentin)",
    },
]


def resolve_path(rel_or_glob: dict) -> list[Path]:
    if rel_or_glob.get("paths"):
        return [TP_RECUP / p for p in rel_or_glob["paths"]]
    return list(TP_RECUP.glob(rel_or_glob["glob"]))


def content_from_path(p: Path) -> dict:
    """Construit un content dict comme extract_content, depuis un fichier/dossier local."""
    if p.is_dir():
        # zipper virtuellement le dossier
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in p.rglob("*"):
                if f.is_file():
                    zf.write(f, arcname=str(f.relative_to(p)).replace("\\", "/"))
        data = buf.getvalue()
        zinfo = b.analyze_zip_bytes(data)
        return {
            "filename": p.name + "/",
            "size": sum(f.stat().st_size for f in p.rglob("*") if f.is_file()),
            "text": p.name + "\n" + zinfo["text"],
            "images": zinfo["images"],
            "videos": zinfo.get("videos", 0),
            "inos": zinfo.get("inos", []),
            "ext": "dir",
        }
    name = p.name.lower()
    data = p.read_bytes()
    out = {
        "filename": p.name,
        "size": p.stat().st_size,
        "text": p.name + "\n",
        "images": 0,
        "videos": 0,
        "inos": [],
        "ext": p.suffix.lower(),
    }
    if name.endswith(".docx"):
        d = b.analyze_docx_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".odt"):
        d = b.analyze_odt_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".pdf"):
        out["text"] += b.read_pdf_text(data)
    elif name.endswith(".zip"):
        zinfo = b.analyze_zip_bytes(data)
        out["text"] += zinfo["text"]
        out["images"] = zinfo["images"]
        out["videos"] = zinfo.get("videos", 0)
        out["inos"] = zinfo["inos"]
    elif name.endswith((".mp4", ".mov", ".m4v", ".avi")):
        out["videos"] = 1
        out["text"] += f"\n[video-file:{p.name}]"
    else:
        out["text"] += data.decode("utf-8", errors="ignore")
    return out


def grade_assignment(asg: dict) -> dict:
    paths = resolve_path(asg)
    missing = [str(p) for p in paths if not p.exists()]
    if not paths or missing:
        raise FileNotFoundError(f"{asg['label']}: missing {missing or 'glob empty'}")
    contents = [content_from_path(p) for p in paths]
    merged = b.merge_contents(contents)
    s2 = b.score_tp2(merged)
    tot = sum(v for k, v in s2.items() if not k.startswith("_"))
    return {
        "label": asg["label"],
        "keys": [asg["key"]] + asg.get("also", []),
        "filename": " | ".join(c["filename"] for c in contents),
        "scores": s2,
        "total": tot,
        "images": merged["images"],
        "videos": merged["videos"],
        "text_preview": merged["text"][:1500],
    }


def apply_tp2_to_workbook(wb, graded: list[dict]):
    """Met à jour les colonnes TP2 pour les élèves concernés (feuilles classes)."""
    # Map clé -> scores
    by_key = {}
    for g in graded:
        for k in g["keys"]:
            by_key[k] = g

    # Colonnes : A=élève B=fichier, TP1 qs puis total, puis TP2
    # TP1_COLS=8 → cols 3..10 qs, 11 total TP1
    # TP2_COLS=13 → cols 12..24 qs, 25 total TP2, 26 grand total, 27 commentaire
    tp1_n = len(b.TP1_COLS)
    tp2_n = len(b.TP2_COLS)
    col_tp2_first = 3 + tp1_n + 1  # 12
    col_tp2_tot = col_tp2_first + tp2_n  # 25
    col_grand = col_tp2_tot + 1  # 26
    col_comment = col_grand + 1  # 27
    col_tp1_tot = 3 + tp1_n  # 11

    for sheet_name in wb.sheetnames:
        if sheet_name == "Barème":
            continue
        ws = wb[sheet_name]
        for r in range(3, ws.max_row + 1):
            nom_cell = (ws.cell(r, 1).value or "").strip()
            if not nom_cell:
                continue
            # retrouver clé
            parts = nom_cell.split()
            if len(parts) < 2:
                continue
            # nom peut être composé
            # Essayer match via eleves
            matched = None
            for key in by_key:
                # key = NOM_PRENOM
                kn = key.replace("_", " ")
                # normalize accents roughly
                def norm(s):
                    s = s.upper()
                    for a, c in [
                        ("É", "E"), ("È", "E"), ("Ê", "E"), ("Ë", "E"),
                        ("À", "A"), ("Â", "A"), ("Ù", "U"), ("Û", "U"),
                        ("Ô", "O"), ("Î", "I"), ("Ï", "I"), ("Ç", "C"), ("Œ", "OE"),
                    ]:
                        s = s.replace(a, c)
                    return s

                if norm(nom_cell).replace(" ", "_") == key or norm(nom_cell).startswith(
                    key.split("_")[0]
                ) and key.split("_")[-1][:4] in norm(nom_cell):
                    # finer: last token prenom
                    prenom = parts[-1]
                    nom = " ".join(parts[:-1])
                    if b.norm_key(nom, prenom) == key or b.norm_key(nom, prenom.split("-")[0]) == key:
                        matched = key
                        break
                    # try all eleves-style
                    if norm(prenom)[:4] in key and norm(nom).replace(" ", "_").replace("-", "-")[:6] in key:
                        matched = key
                        break

            if not matched:
                # direct norm_key
                prenom = parts[-1]
                nom = " ".join(parts[:-1])
                nk = b.norm_key(nom, prenom)
                if nk in by_key:
                    matched = nk
                else:
                    # fuzzy startswith nom
                    for key in by_key:
                        if key.startswith(b.norm_key(nom, "").rstrip("_")) and b.norm_key("", prenom).strip("_")[:3] in key:
                            matched = key
                            break

            if not matched:
                continue

            g = by_key[matched]
            s2 = g["scores"]
            if not s2.get("_rendu"):
                continue

            # append filename note
            prev_f = ws.cell(r, 2).value or ""
            add = f" + TP2:{g['label']}"
            if "TP2:" not in str(prev_f):
                ws.cell(r, 2, str(prev_f) + add)

            for i, (k, mx, _lab) in enumerate(b.TP2_COLS):
                ws.cell(r, col_tp2_first + i, s2.get(k, 0))
            tot2 = g["total"]
            ws.cell(r, col_tp2_tot, tot2)
            ws.cell(r, col_tp2_tot).font = Font(bold=True)

            tot1 = ws.cell(r, col_tp1_tot).value
            try:
                t1 = float(tot1) if tot1 not in (None, "") else 0
            except Exception:
                t1 = 0
            ws.cell(r, col_grand, t1 + tot2)
            ws.cell(r, col_grand).font = Font(bold=True)

            com = ws.cell(r, col_comment).value or ""
            com = str(com).replace("TP2 non identifiable / non rendu", "").strip(" ;")
            note = f"TP2 corrigé ({g['label']}) {tot2}/20"
            if note not in com:
                ws.cell(r, col_comment, (com + " ; " if com else "") + note)

            print(f"  OK {sheet_name} | {nom_cell} <- {matched} TP2={tot2}/20")


def main():
    print("TP_RECUP =", TP_RECUP, "exists=", TP_RECUP.exists())
    graded = []
    for asg in ASSIGNMENTS:
        print("\n===", asg["label"], "===")
        g = grade_assignment(asg)
        graded.append(g)
        print("  keys:", g["keys"])
        print("  imgs/vids:", g["images"], g["videos"], "TP2=", g["total"], "/20")
        detail = {k: v for k, v in g["scores"].items() if not k.startswith("_")}
        print("  detail:", detail)
        # dump preview
        prev_path = Path(__file__).resolve().parent / f"preview_{g['keys'][0]}.txt"
        prev_path.write_text(g["text_preview"], encoding="utf-8")

    if not ORIG_XLSX.exists():
        # fallback corrige
        alt = ORIG_XLSX.with_name("Notes_TP1_TP2_1ere_STI2D_corrige.xlsx")
        src = alt if alt.exists() else None
        if not src:
            raise SystemExit(f"Excel original introuvable: {ORIG_XLSX}")
        print("Original manquant, base =", src)
        shutil.copy2(src, NEW_XLSX)
    else:
        shutil.copy2(ORIG_XLSX, NEW_XLSX)
        print("Copie:", ORIG_XLSX.name, "->", NEW_XLSX.name)

    wb = load_workbook(NEW_XLSX)
    apply_tp2_to_workbook(wb, graded)
    # feuille récap TP2
    if "TP2 ajoutés" in wb.sheetnames:
        del wb["TP2 ajoutés"]
    ws = wb.create_sheet("TP2 ajoutés", 0)
    ws["A1"] = "TP2 ajoutés (sans modifier l'original Notes_TP1_TP2_1ere_STI2D.xlsx)"
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Élèves", "Fichier source", "Total TP2 /20", "Détail"])
    for g in graded:
        detail = ", ".join(f"{k}={v}" for k, v in g["scores"].items() if not k.startswith("_") and v)
        ws.append([", ".join(g["keys"]), g["label"], g["total"], detail])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 80

    wb.save(NEW_XLSX)
    print("\nExcel écrit (nouveau fichier):", NEW_XLSX)
    print("Original intact:", ORIG_XLSX)


if __name__ == "__main__":
    main()
