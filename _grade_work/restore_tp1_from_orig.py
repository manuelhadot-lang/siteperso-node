# -*- coding: utf-8 -*-
"""Restaure les colonnes TP1 depuis l'original si la maj les a dégradées."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
import build_notes_excel as b

ORIG = Path(r"..\doc\Veilleur_intelligent\Notes_TP1_TP2_1ere_STI2D.xlsx")
MAJ = Path(r"..\doc\Veilleur_intelligent\Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx")

wb_o = load_workbook(ORIG, data_only=False)
wb_m = load_workbook(MAJ)

tp1_n = len(b.TP1_COLS)
tp2_n = len(b.TP2_COLS)
col_tp1_tot = 3 + tp1_n
col_tp2_tot = col_tp1_tot + 1 + tp2_n
col_grand = col_tp2_tot + 1


def tp1_total(ws, row):
    vals = []
    for c in range(3, col_tp1_tot):
        v = ws.cell(row, c).value
        if isinstance(v, (int, float)):
            vals.append(v)
    tot = ws.cell(row, col_tp1_tot).value
    if isinstance(tot, (int, float)):
        return tot, True
    if vals:
        return sum(vals), True
    return 0, False


for sheet in wb_m.sheetnames:
    if sheet not in wb_o.sheetnames:
        continue
    if sheet.startswith("Ajouts") or sheet.startswith("TP2") or sheet == "Barème":
        continue
    wo, wm = wb_o[sheet], wb_m[sheet]
    for r in range(3, wm.max_row + 1):
        name = wm.cell(r, 1).value
        if not name:
            continue
        # find same name in original
        ro = None
        for rr in range(3, wo.max_row + 1):
            if wo.cell(rr, 1).value == name:
                ro = rr
                break
        if not ro:
            continue
        old_tot, old_ok = tp1_total(wo, ro)
        new_tot, new_ok = tp1_total(wm, r)
        # restore if original had better/real TP1 and maj lost it or degraded
        if old_ok and old_tot > 0 and (not new_ok or new_tot < old_tot - 0.5 or isinstance(wm.cell(r, col_tp1_tot).value, str)):
            for c in range(3, col_tp1_tot + 1):
                wm.cell(r, c).value = wo.cell(ro, c).value
            # recalc grand
            try:
                t1 = float(wm.cell(r, col_tp1_tot).value or 0)
            except Exception:
                t1 = sum(
                    float(wm.cell(r, c).value or 0)
                    for c in range(3, col_tp1_tot)
                    if isinstance(wm.cell(r, c).value, (int, float))
                )
            try:
                t2 = float(wm.cell(r, col_tp2_tot).value or 0)
            except Exception:
                t2 = 0
            if wm.cell(r, col_tp1_tot).value not in (None, "") or wm.cell(r, col_tp2_tot).value not in (None, ""):
                wm.cell(r, col_grand, t1 + t2)
                wm.cell(r, col_grand).font = Font(bold=True)
            print(f"RESTORE TP1 {sheet} {name}: {new_tot} -> {old_tot} ; grand={t1}+{t2}")

wb_m.save(MAJ)
print("saved", MAJ)
