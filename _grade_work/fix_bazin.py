# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
import build_notes_excel as b

p = Path(r"..\doc\Veilleur_intelligent\Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx")
wb = load_workbook(p)
ws = wb["1ère STI2D 2"]

tp1_n = len(b.TP1_COLS)
tp2_n = len(b.TP2_COLS)
col_tp2 = 3 + tp1_n + 1
col_tp2_tot = col_tp2 + tp2_n
col_grand = col_tp2_tot + 1
col_com = col_grand + 1

# scores from CHEVANNE (same zip)
chevanne = {}
for r in range(3, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or "")
    if "CHEVANNE" in name.upper():
        for i, (k, mx, _) in enumerate(b.TP2_COLS):
            chevanne[k] = ws.cell(r, col_tp2 + i).value or 0
        print("CHEVANNE tot", sum(chevanne.values()))

for r in range(3, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or "")
    if name.upper().startswith("BAZIN"):
        t2 = 0
        for i, (k, mx, _) in enumerate(b.TP2_COLS):
            v = chevanne.get(k, 0)
            ws.cell(r, col_tp2 + i, v)
            t2 += v
        ws.cell(r, col_tp2_tot, t2)
        ws.cell(r, col_tp2_tot).font = Font(bold=True)
        try:
            t1 = float(ws.cell(r, 3 + tp1_n).value or 0)
        except Exception:
            t1 = 0
        ws.cell(r, col_grand, t1 + t2)
        ws.cell(r, col_grand).font = Font(bold=True)
        com = str(ws.cell(r, col_com).value or "")
        com = com.replace("TP2=8/20 (binome Lemonde-San TP2 images)", "")
        com = com.replace("TP2=8/20 (relecture images schema/sim)", "")
        note = "TP2=18/20 (zip Bazin+Chevanne restaure)"
        if note not in com:
            ws.cell(r, col_com, (com.strip(" ;") + " ; " if com.strip(" ;") else "") + note)
        print("BAZIN restored", t2)

wb.save(p)
print("saved")
