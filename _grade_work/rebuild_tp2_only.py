# -*- coding: utf-8 -*-
"""
Remplit correctement les notes TP2 + commentaires dans Notes_TP1_TP2.xlsx
SANS modifier aucune cellule TP1 (colonnes B–J).
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "doc" / "Veilleur_intelligent" / "Notes_TP1_TP2.xlsx"
BACKUP = ROOT / "doc" / "Veilleur_intelligent" / "Notes_TP1_TP2_avant_rebuild_TP2.xlsx"

# Clés barème TP2 (ordre = colonnes K→W)
TP2_KEYS = [
    ("TP2_Q1_Schema", 1, "Schéma pont"),
    ("TP2_Q2_Formule", 1, "Formule"),
    ("TP2_Q3_Egal", 2, "R1=R2"),
    ("TP2_Q4_LDR", 1, "Signif. LDR"),
    ("TP2_Q5_Role", 1, "Rôle LDR"),
    ("TP2_Q6_Multi", 1, "Multimètre"),
    ("TP2_Q7_Nbin", 2, "N déc/bin"),
    ("TP2_Q8_Spice", 1, "Spice"),
    ("TP2_Q9_Tableau", 3, "Tableau mesures"),
    ("TP2_Q10_VideoU", 2, "Vidéo tension"),
    ("TP2_Q11_Comment", 2, "Commenter code"),
    ("TP2_Q12_SimLx", 2, "Sim. lx"),
    ("TP2_Q13_Concl", 1, "Conclusion"),
]

# Notes TP2 consolidées (historique corrections) — détail par question
# Format: scores dict + source + commentaire libre
TP2_GRADES: dict[str, dict] = {
    # —— STI2D1 ——
    "ANCENAY Romain": {
        "src": "SIN_TP2 Benoit+Ancenay",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=0, TP2_Q4_LDR=1, TP2_Q5_Role=0,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "BENOIT Constance": {
        "src": "SIN_TP2 Benoit+Ancenay",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=0, TP2_Q4_LDR=1, TP2_Q5_Role=0,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "BOUCHER Jules": {
        "src": "Boucher/Tp2.docx",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=1, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=0,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "COCOL Alexis": {
        "src": "TP2 LDR Alexis+Corentin",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=0, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "DIOWO Corentin": {
        "src": "TP2 LDR Alexis+Corentin",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=0, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "HOEFFLER Léo": {
        "src": "TP2 Hoeffler+Wipff",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "WIPFF Hugo": {
        "src": "TP2 Hoeffler+Wipff",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "KOMMEGNE Annaëlle": {
        "src": "KOMMEGNE 2026-10-12.odt",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=0, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=0,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "KUCHLY Lucien": {
        "src": "KUCHLY TP2 + vidéo (binôme Léger)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=1, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "LEGER Maxime": {
        "src": "LEGER TP2 + vidéo (binôme Kuchly)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=1, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "MEHREZI Nahil": {
        "src": "MEHREZI TP2 + vidéo",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "MOREAU Antoine": {
        "src": "MOREAU/RESTOUT TP2 (backup)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=0,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "RESTOUT Raphaël": {
        "src": "MOREAU/RESTOUT TP2 (backup)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=0,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    # —— STI2D2 ——
    "BAZIN Francois": {
        "src": "BAZIN+CHEVANNE 13/10",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "CHEVANNE-CARENCO Gabriel": {
        "src": "BAZIN+CHEVANNE 13/10",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "COJANU Alexandre": {
        "src": "COJANU+MEUDIC 13/10",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=0, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "MEUDIC Thomas": {
        "src": "COJANU+MEUDIC 13/10",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=0, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "D'HOKERS-ALLUSSON Adam": {
        "src": "D'Hokers odt+vidéo 13/10",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "FEAUGAS Jean Charles": {
        "src": "binôme D'Hokers (même odt) + vidéo",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "FRANCOIS Alban": {
        # meilleur rendu = Normand_Francois (18) plutôt que Lemonde images (8)
        "src": "Normand_Francois (TP2.docx + vidéo)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "NORMAND Maxime": {
        "src": "Normand_Francois (TP2.docx + vidéo)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "GROULT Grégoire": {
        "src": "GROULT/NOEL zip TP2",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "NOEL Timothée": {
        "src": "GROULT/NOEL zip TP2",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "LEMONDE-SAN Sorya": {
        "src": "Lemonde+François images schéma/sim (relecture visuelle)",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=0, TP2_Q3_Egal=0, TP2_Q4_LDR=1, TP2_Q5_Role=0,
                  TP2_Q6_Multi=0, TP2_Q7_Nbin=0, TP2_Q8_Spice=1, TP2_Q9_Tableau=2, TP2_Q10_VideoU=0,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "LEYE Mariama": {
        "src": "LEYE 13/10.pdf",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "RACHI Marwane": {
        "src": "RACHI 13/10.docx",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=0, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "RECHE Victor": {
        "src": "RECHE 13/10.pdf",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=0, TP2_Q5_Role=0,
                  TP2_Q6_Multi=0, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=0, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
    "RIGAMONTI Clément": {
        "src": "RIGAMONTI 13/10.pdf",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=2, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=0),
    },
    "VARILLON Amaury": {
        "src": "VARILLON 13/10 zip + vidéo",
        "s": dict(TP2_Q1_Schema=1, TP2_Q2_Formule=1, TP2_Q3_Egal=1, TP2_Q4_LDR=1, TP2_Q5_Role=1,
                  TP2_Q6_Multi=1, TP2_Q7_Nbin=2, TP2_Q8_Spice=0, TP2_Q9_Tableau=3, TP2_Q10_VideoU=2,
                  TP2_Q11_Comment=2, TP2_Q12_SimLx=2, TP2_Q13_Concl=1),
    },
}

# Élèves sans TP2 identifiable (commentaire uniquement)
NO_TP2 = {
    "BIENFAIT Noé", "CAROL Thibaut", "DARMON FONSAGRIVES Jules", "DUPONT Foucauld",
    "MONTERO Pierre", "ROUSSEAU Pierre", "DE LAMARZERELLE Théophile", "LOCHEGNIES Jules",
    "SIMON Arthur", "BODILIS Paul", "GIRCOURT Mathieu", "BRAHMI Kalis", "PROSPERI Flora",
    "ROHAUT Andréa", "ROUVEL Gabin", "VERDYS Maxime", "PERE Alban",
}


def norm_name(s: str) -> str:
    s = (s or "").strip().upper()
    for a, c in [
        ("É", "E"), ("È", "E"), ("Ê", "E"), ("Ë", "E"), ("À", "A"), ("Â", "A"),
        ("Ù", "U"), ("Û", "U"), ("Ô", "O"), ("Î", "I"), ("Ï", "I"), ("Ç", "C"),
        ("Œ", "OE"), ("’", "'"),
    ]:
        s = s.replace(a, c)
    return re.sub(r"\s+", " ", s)


def find_grade(student: str) -> dict | None:
    n = norm_name(student)
    for k, g in TP2_GRADES.items():
        if norm_name(k) == n:
            return g
    # fuzzy: same family name start
    for k, g in TP2_GRADES.items():
        kn = norm_name(k)
        if kn.split()[0] == n.split()[0] and kn.split()[-1][:4] == n.split()[-1][:4]:
            return g
    return None


def build_comment(student: str, grade: dict | None) -> str:
    if grade is None:
        if any(norm_name(student) == norm_name(x) for x in NO_TP2):
            return "TP2 : non rendu / non identifiable dans les dépôts analysés."
        return "TP2 : non corrigé (pas de copie associée)."
    s = grade["s"]
    tot = sum(s.get(k, 0) for k, _, _ in TP2_KEYS)
    missing = []
    for k, mx, lab in TP2_KEYS:
        v = s.get(k, 0)
        if v < mx:
            missing.append(f"{lab} {v}/{mx}")
    parts = [f"TP2 {tot}/20 — source : {grade['src']}."]
    if missing:
        parts.append("Manques : " + "; ".join(missing) + ".")
    else:
        parts.append("Barème complet.")
    return " ".join(parts)


def main():
    if not XLSX.exists():
        raise SystemExit(f"Fichier introuvable: {XLSX}")
    shutil.copy2(XLSX, BACKUP)
    print("Backup:", BACKUP.name)

    wb = load_workbook(XLSX)
    # colonnes: A=1 élève, B-I TP1, J=10 tot TP1, K-W=11-23 TP2 qs, X=24 tot TP2, Y=25 commentaire
    COL_TP2_FIRST = 11  # K
    COL_TP2_TOT = 24    # X
    COL_COMMENT = 25    # Y

    updated = 0
    cleared = 0
    for sheet_name in ("1ère STI2D 1", "1ère STI2D 2"):
        if sheet_name not in wb.sheetnames:
            print("Feuille manquante:", sheet_name)
            continue
        ws = wb[sheet_name]
        print("\n===", sheet_name, "===")
        for r in range(3, ws.max_row + 1):
            student = ws.cell(r, 1).value
            if not student:
                continue
            # Snapshot TP1 pour vérifier qu'on ne touche pas
            tp1_before = [ws.cell(r, c).value for c in range(2, 11)]

            grade = find_grade(str(student))
            if grade:
                tot = 0
                for i, (k, mx, _) in enumerate(TP2_KEYS):
                    v = int(grade["s"].get(k, 0))
                    ws.cell(r, COL_TP2_FIRST + i, v)
                    tot += v
                ws.cell(r, COL_TP2_TOT, tot)
                ws.cell(r, COL_TP2_TOT).font = Font(bold=True)
                ws.cell(r, COL_COMMENT, build_comment(str(student), grade))
                ws.cell(r, COL_COMMENT).alignment = Alignment(wrap_text=True)
                updated += 1
                print(f"  {student}: TP2={tot}/20")
            else:
                # vider TP2 erroné éventuel + commentaire clair
                # openpyxl: cell(r,c,None) N'EFFACE PAS — il faut .value = None
                for i in range(len(TP2_KEYS)):
                    ws.cell(r, COL_TP2_FIRST + i).value = None
                ws.cell(r, COL_TP2_TOT).value = None
                ws.cell(r, COL_COMMENT).value = build_comment(str(student), None)
                cleared += 1
                print(f"  {student}: TP2 vide (non rendu)")

            tp1_after = [ws.cell(r, c).value for c in range(2, 11)]
            if tp1_before != tp1_after:
                raise RuntimeError(f"BUG: TP1 modifié pour {student} !")

    wb.save(XLSX)
    print(f"\nOK écrit: {XLSX}")
    print(f"TP2 remplis: {updated} | sans TP2: {cleared}")
    print("TP1 non modifié. Backup:", BACKUP)


if __name__ == "__main__":
    main()
