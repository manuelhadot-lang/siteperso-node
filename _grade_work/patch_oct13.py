# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
import build_notes_excel as b

p = Path(r"..\doc\Veilleur_intelligent\Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx")
wb = load_workbook(p)

LEMONDE = {
    "TP2_Q1_Schema": 1,
    "TP2_Q2_Formule": 0,
    "TP2_Q3_Egal": 0,
    "TP2_Q4_LDR": 1,
    "TP2_Q5_Role": 0,
    "TP2_Q6_Multi": 0,
    "TP2_Q7_Nbin": 0,
    "TP2_Q8_Spice": 1,
    "TP2_Q9_Tableau": 2,
    "TP2_Q10_VideoU": 0,
    "TP2_Q11_Comment": 0,
    "TP2_Q12_SimLx": 2,
    "TP2_Q13_Concl": 1,
}  # 8/20 — schémas + simu images, peu de texte

tp1_n = len(b.TP1_COLS)
tp2_n = len(b.TP2_COLS)
col_tp2 = 3 + tp1_n + 1
col_tp2_tot = col_tp2 + tp2_n
col_grand = col_tp2_tot + 1
col_com = col_grand + 1


def write_tp2(ws, row, scores, label):
    t2 = 0
    for i, (k, mx, _) in enumerate(b.TP2_COLS):
        v = scores.get(k, 0) or 0
        ws.cell(row, col_tp2 + i, v)
        t2 += v
    ws.cell(row, col_tp2_tot, t2)
    ws.cell(row, col_tp2_tot).font = Font(bold=True)
    try:
        t1 = float(ws.cell(row, 3 + tp1_n).value or 0)
    except Exception:
        t1 = 0
    ws.cell(row, col_grand, t1 + t2)
    ws.cell(row, col_grand).font = Font(bold=True)
    f = str(ws.cell(row, 2).value or "")
    if label not in f:
        ws.cell(row, 2, f + " + " + label)
    com = str(ws.cell(row, col_com).value or "")
    note = f"TP2={t2}/20 ({label})"
    if note not in com:
        ws.cell(row, col_com, (com + " ; " if com else "") + note)
    print(ws.cell(row, 1).value, "TP2", t2)


ws2 = wb["1ère STI2D 2"]
dhokers = {}
for r in range(3, ws2.max_row + 1):
    name = str(ws2.cell(r, 1).value or "")
    up = name.upper()
    if "HOKERS" in up:
        for i, (k, mx, _) in enumerate(b.TP2_COLS):
            dhokers[k] = ws2.cell(r, col_tp2 + i).value or 0
        print("DHokers", sum(dhokers.values()), dhokers)

for r in range(3, ws2.max_row + 1):
    name = str(ws2.cell(r, 1).value or "")
    up = name.upper()
    if "LEMONDE" in up:
        write_tp2(ws2, r, LEMONDE, "relecture images schema/sim")
    if "FRANCOIS" in up:
        write_tp2(ws2, r, LEMONDE, "binome Lemonde-San TP2 images")
    if "FEAUGAS" in up and dhokers:
        write_tp2(ws2, r, dhokers, "binome DHokers + video")

if "Ajouts 13 oct" in wb.sheetnames:
    ws = wb["Ajouts 13 oct"]
    ws.append(["LEMONDE-SAN + FRANCOIS (images)", "zip Sorya+Alban", "—", 8])
    ws.append(["FEAUGAS (= binome DHokers)", "odt + video Adam", "—", sum(dhokers.values())])

wb.save(p)
print("saved", p)
