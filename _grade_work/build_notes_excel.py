# -*- coding: utf-8 -*-
"""
Construit un Excel de notes TP1/TP2 Veilleur Intelligent
à partir de eleves.json + copies (local + backup ZIP).
"""
from __future__ import annotations

import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WORK = Path(__file__).resolve().parent
LOCAL_UPLOAD = ROOT / "upload-tp" / "Veilleur_intelligent"
BACKUP_ZIP = Path(r"c:\Users\mhado\Downloads\BACKUP_STI2D_2026-10-11.zip")
OUT_XLSX = ROOT / "doc" / "Veilleur_intelligent" / "Notes_TP1_TP2_1ere_STI2D.xlsx"

# Barème exact TP1_LED.pdf (total /20)
TP1_COLS = [
    ("TP1_Q1_SimVide", 1, "Sim. Spice — programme vide ?"),
    ("TP1_Q2_Allumer", 1, "Que faire pour allumer la LED ?"),
    ("TP1_Q3_ProgOn", 3, "Prog. allumer LED (+ photo sim.)"),
    ("TP1_Q4_ProgBlink", 4, "Prog. clignotement 500 ms (+ photo sim.)"),
    ("TP1_Q5_PhotoMontage", 4, "Photo du câblage / montage"),
    ("TP1_Q6_VideoCligno", 3, "Vidéo LED clignote ≈1/s"),
    ("TP1_Q7_VideoDelais", 2, "Vidéo délais 200 ms / 1000 ms"),
    ("TP1_Q8_Inverse", 2, "Option inverse LOW/HIGH + conclusion"),
]
# Barème exact TP2_LDR_mesure.pdf (total /20)
TP2_COLS = [
    ("TP2_Q1_Schema", 1, "Schéma pont diviseur (E, R1, R2, V)"),
    ("TP2_Q2_Formule", 1, "Formule du pont diviseur"),
    ("TP2_Q3_Egal", 2, "Si R1=R2, que vaut V ?"),
    ("TP2_Q4_LDR", 1, "Signification des lettres LDR"),
    ("TP2_Q5_Role", 1, "Rôle de la LDR dans le projet"),
    ("TP2_Q6_Multi", 1, "Que mesure un multimètre ?"),
    ("TP2_Q7_Nbin", 2, "N décimal/binaire (0 / 1,6 / 3,3 V)"),
    ("TP2_Q8_Spice", 1, "Schéma saisi sur Simulateur Spice"),
    ("TP2_Q9_Tableau", 3, "Tableau mesures 3 situations"),
    ("TP2_Q10_VideoU", 2, "Vidéo évolution tension / lumière"),
    ("TP2_Q11_Comment", 2, "Commenter le programme ligne à ligne"),
    ("TP2_Q12_SimLx", 2, "ADC/U pour 10000 / 10 / 1 lx"),
    ("TP2_Q13_Concl", 1, "Conclusion"),
]
TP1_TOTAL = sum(m for _, m, _ in TP1_COLS)  # 20
TP2_TOTAL = sum(m for _, m, _ in TP2_COLS)  # 20


def norm_key(nom: str, prenom: str) -> str:
    s = f"{nom}_{prenom}".upper()
    s = (
        s.replace("É", "E")
        .replace("È", "E")
        .replace("Ê", "E")
        .replace("Ë", "E")
        .replace("À", "A")
        .replace("Â", "A")
        .replace("Ä", "A")
        .replace("Ù", "U")
        .replace("Û", "U")
        .replace("Ü", "U")
        .replace("Ô", "O")
        .replace("Ö", "O")
        .replace("Î", "I")
        .replace("Ï", "I")
        .replace("Ç", "C")
        .replace("Œ", "OE")
        .replace("'", "")
        .replace("’", "")
        .replace(" ", "_")
        .replace("-", "-")
    )
    return re.sub(r"[^A-Z0-9_\-]", "", s)


def parse_submission_name(filename: str):
    """NOM_Prenom_YYYY-MM-DD.ext → (key, date, name)"""
    stem = Path(filename).stem
    m = re.match(r"^(.+)_(\d{4}-\d{2}-\d{2})$", stem)
    if not m:
        return None, None, stem
    who, date_s = m.group(1), m.group(2)
    try:
        dt = datetime.strptime(date_s, "%Y-%m-%d")
    except ValueError:
        dt = None
    key = re.sub(r"[^A-Z0-9_\-]", "", who.upper().replace("É", "E").replace("È", "E").replace("Ê", "E")
                 .replace("À", "A").replace("Ù", "U").replace("Ç", "C").replace("Ô", "O")
                 .replace("Î", "I").replace("Ï", "I").replace("Ë", "E").replace("Ü", "U")
                 .replace("Ö", "O").replace("Â", "A").replace("Û", "U"))
    # soften accents already handled roughly
    key = (
        who.upper()
        .replace("É", "E").replace("È", "E").replace("Ê", "E").replace("Ë", "E")
        .replace("À", "A").replace("Â", "A").replace("Ä", "A")
        .replace("Ù", "U").replace("Û", "U").replace("Ü", "U")
        .replace("Ô", "O").replace("Ö", "O")
        .replace("Î", "I").replace("Ï", "I")
        .replace("Ç", "C")
    )
    key = re.sub(r"[^A-Z0-9_\-]", "", key)
    return key, dt, who


def load_eleves():
    path = WORK / "eleves.json"
    return json.loads(path.read_text(encoding="utf-8-sig"))


def collect_submissions():
    """Map student_key -> list of (date, path_or_zipmember, source_label, size)."""
    by_key = defaultdict(list)

    # Local files
    if LOCAL_UPLOAD.exists():
        for p in LOCAL_UPLOAD.iterdir():
            if not p.is_file() or p.name.startswith("."):
                continue
            key, dt, _ = parse_submission_name(p.name)
            if not key:
                continue
            by_key[key].append(
                {
                    "date": dt or datetime.min,
                    "name": p.name,
                    "kind": "local",
                    "path": p,
                    "size": p.stat().st_size,
                }
            )

    # Backup zip members
    if BACKUP_ZIP.exists():
        with zipfile.ZipFile(BACKUP_ZIP, "r") as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                if not info.filename.startswith("upload-tp/Veilleur_intelligent/"):
                    continue
                fname = Path(info.filename).name
                if not fname or fname.startswith("."):
                    continue
                key, dt, _ = parse_submission_name(fname)
                if not key:
                    continue
                by_key[key].append(
                    {
                        "date": dt or datetime.min,
                        "name": fname,
                        "kind": "backup",
                        "zip_name": info.filename,
                        "size": info.file_size,
                    }
                )

    # Keep latest per student (by date then size)
    latest = {}
    for key, items in by_key.items():
        items.sort(key=lambda x: (x["date"], x["size"]), reverse=True)
        latest[key] = items[0]
    return latest, by_key


def _bytes_looks_like_video(raw: bytes) -> bool:
    """Détecte une vidéo embarquée (OLE Word, paquet, etc.)."""
    if not raw:
        return False
    head = raw[:64]
    # signatures conteneurs MP4 / QuickTime / RIFF AVI
    if b"ftyp" in head or b"ftypmp4" in raw[:128] or b"ftypisom" in raw[:128] or b"ftypqt" in raw[:128]:
        return True
    if raw[:4] == b"RIFF" and b"AVI" in raw[:16]:
        return True
    # OLE Word stocke souvent le nom .mp4/.mov dans le binaire
    low = raw[:8000].lower()
    if b".mp4" in low or b".mov" in low or b".m4v" in low or b"video/mp4" in low:
        return True
    # balayage léger du fichier (vidéos OLE ~quelques Mo)
    if len(raw) > 8000 and (b"ftypmp4" in raw or b"ftypisom" in raw or b"mdat" in raw[:200000]):
        # mdat seul est trop large — exiger aussi ftyp quelque part au début
        if b"ftyp" in raw[:500000]:
            return True
    return False


def _extract_urls_from_bytes(raw: bytes) -> list[str]:
    try:
        txt = raw.decode("utf-8", errors="ignore")
    except Exception:
        return []
    return re.findall(r"https?://[^\s\"'<>]+", txt)


def analyze_docx_bytes(data: bytes) -> dict:
    """Texte + images + vidéos OLE/liens hypertexte d'un .docx."""
    info = {"text": "", "images": 0, "videos": 0, "inos": [], "links": []}
    try:
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            # texte visible
            try:
                xml = zf.read("word/document.xml")
                root = ET.fromstring(xml)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                texts = [t.text for t in root.findall(".//w:t", ns) if t.text]
                info["text"] = "\n".join(texts)
                # URLs éventuellement en clair dans le XML
                info["links"].extend(_extract_urls_from_bytes(xml))
            except Exception as e:
                info["text"] = f"[docx-text-error {e}]"

            # hyperliens externes (Target Mode=External)
            for n in zf.namelist():
                if n.endswith(".rels"):
                    try:
                        rel_xml = zf.read(n).decode("utf-8", errors="ignore")
                    except Exception:
                        continue
                    for m in re.finditer(
                        r'Target="(https?://[^"]+)"[^>]*TargetMode="External"|'
                        r'TargetMode="External"[^>]*Target="(https?://[^"]+)"|'
                        r'Target="(https?://[^"]+)"',
                        rel_xml,
                    ):
                        url = next(g for g in m.groups() if g)
                        if "schemas.openxmlformats.org" in url or "schemas.microsoft.com" in url:
                            continue
                        info["links"].append(url)

            # images media
            info["images"] = sum(
                1
                for n in zf.namelist()
                if n.startswith("word/media/")
                and n.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".emf", ".wmf", ".tif"))
            )

            # vidéos : fichiers media vidéo + objets OLE embarqués
            for n in zf.namelist():
                low = n.lower()
                if low.endswith((".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm")):
                    info["videos"] += 1
                    info["text"] += f"\n[video-file:{Path(n).name}]"
                    continue
                if "embeddings/" in low or low.endswith(".bin"):
                    try:
                        raw = zf.read(n)
                    except Exception:
                        continue
                    if _bytes_looks_like_video(raw):
                        info["videos"] += 1
                        # récupérer un nom de fichier mp4 s'il est dans l'OLE
                        m = re.search(rb"([\w.\- ]{1,80}\.(?:mp4|mov|m4v|avi))", raw[:20000], re.I)
                        vname = m.group(1).decode("utf-8", errors="ignore") if m else Path(n).name
                        info["text"] += f"\n[video-ole:{vname}]"
                        info["links"].extend(
                            u
                            for u in _extract_urls_from_bytes(raw[:50000])
                            if not u.startswith("http://schemas.")
                        )

            if info["links"]:
                # dédupliquer en gardant l'ordre
                seen = set()
                uniq = []
                for u in info["links"]:
                    if u not in seen:
                        seen.add(u)
                        uniq.append(u)
                info["links"] = uniq
                info["text"] += "\n" + "\n".join(info["links"])
    except Exception as e:
        info["text"] = f"[docx-error {e}]"
    return info


def read_docx_text(data: bytes) -> str:
    return analyze_docx_bytes(data).get("text", "")


def analyze_odt_bytes(data: bytes) -> dict:
    info = {"text": "", "images": 0, "videos": 0, "inos": [], "links": []}
    try:
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            try:
                xml = zf.read("content.xml")
                root = ET.fromstring(xml)
                texts = []
                for el in root.iter():
                    if el.text:
                        texts.append(el.text)
                    if el.tail:
                        texts.append(el.tail)
                info["text"] = "\n".join(texts)
                info["links"].extend(_extract_urls_from_bytes(xml))
            except Exception as e:
                info["text"] = f"[odt-text-error {e}]"

            # liens dans les styles / meta
            for n in ("meta.xml", "styles.xml", "content.xml"):
                if n in zf.namelist():
                    info["links"].extend(_extract_urls_from_bytes(zf.read(n)))

            info["images"] = sum(
                1 for n in zf.namelist() if "Pictures/" in n or "/media/" in n.lower()
            )
            for n in zf.namelist():
                low = n.lower()
                if low.endswith((".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm")):
                    info["videos"] += 1
                    info["text"] += f"\n[video-file:{Path(n).name}]"
                elif low.endswith((".bin", ".ole")) or "object" in low:
                    try:
                        raw = zf.read(n)
                    except Exception:
                        continue
                    if _bytes_looks_like_video(raw):
                        info["videos"] += 1
                        info["text"] += f"\n[video-ole:{Path(n).name}]"

            if info["links"]:
                seen = set()
                uniq = []
                for u in info["links"]:
                    if u not in seen and "schemas.open" not in u:
                        seen.add(u)
                        uniq.append(u)
                info["links"] = uniq
                info["text"] += "\n" + "\n".join(info["links"])
    except Exception as e:
        info["text"] = f"[odt-error {e}]"
    return info


def read_odt_text(data: bytes) -> str:
    return analyze_odt_bytes(data).get("text", "")


def read_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages[:20]:
            parts.append(page.extract_text() or "")
        # annotations / liens
        for page in reader.pages[:20]:
            if "/Annots" in page:
                try:
                    for annot in page["/Annots"]:
                        obj = annot.get_object()
                        a = obj.get("/A")
                        if a and a.get("/URI"):
                            parts.append(str(a["/URI"]))
                except Exception:
                    pass
        return "\n".join(parts)
    except Exception as e:
        return f"[pdf-error {e}]"


def analyze_zip_bytes(data: bytes) -> dict:
    info = {
        "text": "",
        "inos": [],
        "images": 0,
        "videos": 0,
        "names": [],
    }
    try:
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            for n in zf.namelist():
                info["names"].append(n)
                low = n.lower()
                info["text"] += "\n" + n
                if low.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".heic", ".tif", ".tiff")):
                    info["images"] += 1
                if low.endswith((".mov", ".mp4", ".avi", ".mkv", ".m4v", ".webm")):
                    info["videos"] += 1
                if low.endswith((".ino", ".cpp", ".c", ".h", ".txt", ".md", ".csv", ".json")):
                    try:
                        raw = zf.read(n)
                        txt = raw.decode("utf-8", errors="ignore")
                        if low.endswith(".ino") or "void setup" in txt or "analogRead" in txt:
                            info["inos"].append((n, txt))
                        info["text"] += "\n" + txt
                    except Exception:
                        pass
                if low.endswith(".pdf"):
                    # Ignorer uniquement le sujet officiel vierge (pas les PDF élèves
                    # du type « TP2_LDR_ chevanne+bazin.pdf »)
                    base = Path(n).name.lower().replace(" ", "")
                    official = {
                        "tp1_led.pdf",
                        "tp1_led_esp32.pdf",
                        "tp2_ldr_mesure.pdf",
                        "tp3_final.pdf",
                    }
                    if base in official:
                        info["text"] += f"\n[sujet-ignore:{Path(n).name}]"
                        continue
                    try:
                        info["text"] += "\n" + read_pdf_text(zf.read(n))
                    except Exception:
                        pass
                if low.endswith(".docx"):
                    try:
                        d = analyze_docx_bytes(zf.read(n))
                        info["text"] += "\n" + d["text"]
                        info["images"] += d.get("images", 0)
                        info["videos"] += d.get("videos", 0)
                    except Exception:
                        pass
                if low.endswith(".odt"):
                    try:
                        d = analyze_odt_bytes(zf.read(n))
                        info["text"] += "\n" + d["text"]
                        info["images"] += d.get("images", 0)
                        info["videos"] += d.get("videos", 0)
                    except Exception:
                        pass
    except Exception as e:
        info["text"] = f"[zip-error {e}]"
    return info


def load_submission_bytes(sub: dict) -> bytes:
    if sub["kind"] == "local":
        return Path(sub["path"]).read_bytes()
    with zipfile.ZipFile(BACKUP_ZIP, "r") as zf:
        return zf.read(sub["zip_name"])


def extract_content(sub: dict) -> dict:
    name = sub["name"].lower()
    data = load_submission_bytes(sub)
    out = {
        "filename": sub["name"],
        "size": sub["size"],
        "text": sub["name"] + "\n",
        "images": 0,
        "videos": 0,
        "inos": [],
        "ext": Path(sub["name"]).suffix.lower(),
    }
    if name.endswith(".docx"):
        d = analyze_docx_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".odt"):
        d = analyze_odt_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".pdf"):
        out["text"] += read_pdf_text(data)
    elif name.endswith(".zip"):
        zinfo = analyze_zip_bytes(data)
        out["text"] += zinfo["text"]
        out["images"] = zinfo["images"]
        out["videos"] = zinfo.get("videos", 0)
        out["inos"] = zinfo["inos"]
    else:
        out["text"] += data.decode("utf-8", errors="ignore")
    return out


def merge_contents(contents: list[dict]) -> dict:
    """Fusionne plusieurs dépôts d'un même élève (TP1 + TP2 séparés)."""
    merged = {
        "filename": " + ".join(c["filename"] for c in contents),
        "size": sum(c.get("size", 0) for c in contents),
        "text": "",
        "images": 0,
        "videos": 0,
        "inos": [],
        "ext": "merged",
    }
    for c in contents:
        merged["text"] += "\n" + (c.get("text") or "")
        merged["images"] += c.get("images", 0)
        merged["videos"] += c.get("videos", 0)
        merged["inos"].extend(c.get("inos") or [])
    return merged


def _media_links(blob: str) -> list[str]:
    """Tous les URL susceptibles d'être photo/vidéo (Drive, Immich, ScreenPal…)."""
    return re.findall(r"https?://[^\s)\"'<>]+", blob or "")


def _has_photo_evidence(content: dict, blob: str) -> bool:
    if content.get("images", 0) >= 1:
        return True
    if re.search(r"\.(?:png|jpe?g|gif|webp|heic|bmp)\b", blob):
        return True
    # Lien / dépôt photo explicite (pas la simple consigne « photographier » du sujet)
    if re.search(
        r"lien\s+(pour\s+)?(les\s+)?photos?|mail-attachment|imgur|ibb\.co|"
        r"photo\s+ci[- ]?(jointe|dessous)|voir\s+(la\s+)?photo",
        blob,
    ):
        return True
    for url in _media_links(blob):
        if "schemas.openxmlformats.org" in url or "schemas.microsoft.com" in url:
            continue
        if re.search(
            r"drive\.google|drive\.usercontent|mail-attachment|immich|"
            r"skm-event|/s/photo|/s/esp32|imgur|ibb\.co",
            url,
        ):
            return True
    return False


def _has_video_evidence(content: dict, blob: str) -> bool:
    """Preuve réelle de vidéo (fichier, nom .mp4/.mov, ou lien) — pas le seul texte du sujet."""
    if content.get("videos", 0) >= 1:
        return True
    # fichiers nommés dans une archive / le texte
    if re.search(r"\.(?:mov|mp4|m4v|avi|mkv|webm)\b", blob):
        return True
    if re.search(r"\bvdo\d*\.mp4\b|video_montage_|\.mov\b", blob):
        return True
    # liens vidéo dédiés
    for url in _media_links(blob):
        if re.search(
            r"youtube|youtu\.be|screenpal|vimeo|streamable|"
            r"drive\.google|drive\.usercontent|immich|skm-event|"
            r"/s/200ms|/s/10ms|/s/lowhigh|/s/esp32|/watch/",
            url,
        ):
            return True
    # mention explicite d'un dépôt vidéo (hors simple recopie du barème)
    if re.search(
        r"lien\s+(pour\s+)?(la\s+)?vid[eé]o|vid[eé]os?\s+disponible|"
        r"voir\s+(la\s+)?vid[eé]o|voir\s+vid[eé]o\s+vdo|"
        r"capture\s+vid[eé]o|r[eé]sultat\s+en\s+vid[eé]o|"
        r"vid[eé]o\s+(si\s+dessous|ci-dessous|jointe|dans\s+le\s+fichier|"
        r"du\s+montage|de\s+la\s+led|clignot)",
        blob,
    ):
        return True
    return False


def _count_video_refs(content: dict, blob: str) -> int:
    n = content.get("videos", 0)
    n += len(re.findall(r"\.(?:mov|mp4|m4v|avi|mkv|webm)\b", blob))
    n += len(re.findall(r"\bvdo\d*\.mp4\b", blob))
    for url in _media_links(blob):
        if re.search(
            r"youtube|youtu\.be|screenpal|vimeo|immich|skm-event|/s/\w+|drive\.google|/watch/",
            url,
        ):
            n += 1
    return n


def score_tp1(content: dict) -> dict:
    """Barème TP1_LED.pdf — /20."""
    text = (content.get("text") or "").lower()
    ino_blob = "\n".join(t for _, t in content.get("inos", [])).lower()
    blob = text + "\n" + ino_blob
    photos = _has_photo_evidence(content, blob)
    videos = _has_video_evidence(content, blob)
    n_vid = _count_video_refs(content, blob)

    scores = {}

    # Q1 (1) — programme vide → rien / LED éteinte
    if re.search(
        r"(ne se passe rien|rien ne se passe|led .+ [eé]teinte|pas de courant|"
        r"programme est vide.{0,80}rien|vide.{0,40}(rien|off|[eé]teint))",
        blob,
    ):
        scores["TP1_Q1_SimVide"] = 1
    elif re.search(r"question\s*1|programme est vide|simuler ce montage", blob) and re.search(
        r"rien|[eé]teint|aucun|pas de", blob
    ):
        scores["TP1_Q1_SimVide"] = 1
    else:
        scores["TP1_Q1_SimVide"] = 0

    # Q2 (1) — écrire un programme / configurer sortie HIGH/LOW
    if re.search(
        r"(faut .{0,20}(programme|pinmode|digitalwrite|configurer)|"
        r"[eé]crire un programme|mettre .{0,20}(high|low|sortie)|"
        r"broche .{0,30}sortie)",
        blob,
    ):
        scores["TP1_Q2_Allumer"] = 1
    else:
        scores["TP1_Q2_Allumer"] = 0

    # Q3 (3) — programme allumer LED (pinMode + digitalWrite) ; photo sim attendue
    q3 = 0
    has_on_code = ("pinmode" in blob and "digitalwrite" in blob) or re.search(
        r"digitalwrite\s*\([^)]*(high|low)", blob
    )
    if has_on_code:
        q3 += 2
        # photo simulation ou image jointe
        if photos or re.search(r"spice|simulation|simulateur", blob):
            q3 += 1
        elif "void setup" in blob and "void loop" in blob:
            q3 += 1  # code complet même sans photo explicite → 3/3 si structure OK
            q3 = min(3, q3)
    scores["TP1_Q3_ProgOn"] = min(3, q3)

    # Q4 (4) — clignotement delay(500) ; photo sim
    q4 = 0
    blink = (
        "digitalwrite" in blob
        and "delay" in blob
        and re.search(r"delay\s*\(\s*500\s*\)|500\s*ms|clignot", blob)
    )
    if blink or (re.search(r"delay\s*\(\s*500\s*\)", blob) and "digitalwrite" in blob):
        q4 += 3
        if photos or re.search(r"spice|simulation|photo", blob) or content.get("images", 0) >= 1:
            q4 += 1
    elif "digitalwrite" in blob and "delay" in blob:
        q4 += 2  # clignote mais délai non 500 explicite
    scores["TP1_Q4_ProgBlink"] = min(4, q4)

    # Q5 (4) — photo câblage montage (fichier image OU lien photo)
    if photos:
        scores["TP1_Q5_PhotoMontage"] = 4
    elif re.search(r"photo\s+du\s+montage|photograph", blob):
        scores["TP1_Q5_PhotoMontage"] = 2
    else:
        scores["TP1_Q5_PhotoMontage"] = 0

    # Q6 (3) — vidéo clignotement ~1/s : points dès qu'une vidéo/lien vidéo est présent
    if videos or n_vid >= 1:
        scores["TP1_Q6_VideoCligno"] = 3
    else:
        scores["TP1_Q6_VideoCligno"] = 0

    # Q7 (2) — vidéo 200 ms / 1000 ms
    has_200 = bool(
        re.search(
            r"200\s*ms|delay\s*\(\s*200\s*\)|delay\s*200|_200|/s/200ms|video_montage_200",
            blob,
        )
    )
    has_1000 = bool(
        re.search(
            r"1000\s*ms|delay\s*\(\s*1000\s*\)|delay\s*1000|_1000|"
            r"video_montage_1000|/s/10ms",  # lien Immich « 10ms » souvent utilisé pour délai modifié
            blob,
        )
    )
    # Plusieurs vidéos (ou liens) → on considère les essais de délais fournis
    if videos or n_vid >= 1:
        if has_200 or has_1000 or n_vid >= 2:
            scores["TP1_Q7_VideoDelais"] = 2
        else:
            # une seule vidéo présente : crédit partiel (souvent le cligno 500 ms déjà noté en Q6)
            scores["TP1_Q7_VideoDelais"] = 1
    else:
        scores["TP1_Q7_VideoDelais"] = 0

    # Q8 (2) — inverse LOW/HIGH + conclusion
    if re.search(r"invers|LOW\s*/\s*HIGH|high\s*.{0,20}low|sens du montage|cathode|logique", blob) and re.search(
        r"conclus|donc|allume|rien|toujours|s[' ]allume|[eé]teint", blob
    ):
        scores["TP1_Q8_Inverse"] = 2
    elif re.search(r"invers.*(low|high)|(low|high).*invers", blob):
        scores["TP1_Q8_Inverse"] = 1
    else:
        scores["TP1_Q8_Inverse"] = 0

    tp1_signals = (
        scores["TP1_Q1_SimVide"]
        or scores["TP1_Q2_Allumer"]
        or scores["TP1_Q3_ProgOn"]
        or scores["TP1_Q4_ProgBlink"]
        or scores["TP1_Q5_PhotoMontage"]
        or scores["TP1_Q6_VideoCligno"]
        or "pinmode" in blob
        or "digitalwrite" in blob
        or re.search(r"\btp1\b|allumer une led|clignot", blob)
    )
    scores["_rendu"] = bool(tp1_signals)
    if not scores["_rendu"]:
        for k in list(scores):
            if not k.startswith("_"):
                scores[k] = 0
    return scores


def score_tp2(content: dict) -> dict:
    """Barème TP2_LDR_mesure.pdf — /20."""
    text = (content.get("text") or "").lower()
    ino_blob = "\n".join(t for _, t in content.get("inos", [])).lower()
    blob = text + "\n" + ino_blob
    photos = _has_photo_evidence(content, blob)
    videos = _has_video_evidence(content, blob)

    scores = {}

    # Q1 schéma pont
    if re.search(
        r"pont\s+diviseur|sch[eé]ma.{0,40}(ldr|photor|r1\s*.{0,10}r2)|"
        r"e\s*=\s*3[,.]?3.{0,40}r[12]|v22|point\s+milieu",
        blob,
    ) or (photos and re.search(r"pont\s+diviseur|ldr|photor", blob)):
        scores["TP2_Q1_Schema"] = 1
    else:
        scores["TP2_Q1_Schema"] = 0

    # Q2 formule V = E * R2/(R1+R2) ou équivalent
    if re.search(
        r"r2\s*/\s*\(\s*r1\s*\+\s*r2\s*\)|"
        r"v\s*=\s*e\s*[×x*.]|u\s*=\s*.{0,10}r2|"
        r"formule.{0,40}diviseur|tension.{0,20}divis",
        blob,
    ):
        scores["TP2_Q2_Formule"] = 1
    else:
        scores["TP2_Q2_Formule"] = 0

    # Q3 Si R1=R2 → V = E/2 = 1,65 V
    if re.search(
        r"1[,.]65|e\s*/\s*2|moiti[eé]|3[,.]?3\s*/\s*2|r1\s*=\s*r2.{0,40}(1[,.]65|moiti)",
        blob,
    ):
        scores["TP2_Q3_Egal"] = 2
    elif re.search(r"r1\s*=\s*r2", blob) and re.search(r"1[,.]6|moiti|égal|egale", blob):
        scores["TP2_Q3_Egal"] = 1
    else:
        scores["TP2_Q3_Egal"] = 0

    # Q4 LDR = Light Dependent Resistor / photorésistance
    if re.search(
        r"light\s*dependent|photor[eé]sistance|d[eé]pendante?\s+de\s+la\s+lumi|"
        r"r[eé]sistance.{0,20}lumi",
        blob,
    ):
        scores["TP2_Q4_LDR"] = 1
    else:
        scores["TP2_Q4_LDR"] = 0

    # Q5 rôle dans le projet
    if re.search(
        r"(r[oô]le|sert|permet|capteur).{0,60}(lumi|nuit|jour|obscurit)|"
        r"(mesurer|d[eé]tecter).{0,40}lumi",
        blob,
    ):
        scores["TP2_Q5_Role"] = 1
    else:
        scores["TP2_Q5_Role"] = 0

    # Q6 multimètre
    if re.search(
        r"multim[eè]tre.{0,80}(tension|courant|r[eé]sistance|volt)|"
        r"(mesure|mesurer).{0,40}(tension|volt|courant|ohm)",
        blob,
    ):
        scores["TP2_Q6_Multi"] = 1
    else:
        scores["TP2_Q6_Multi"] = 0

    # Q7 N pour 0 / 1.6 / 3.3 V (0, ~1984, 4095) + binaire
    has_0 = bool(re.search(r"\b0\b.{0,30}(0\s*v|u22.{0,10}0)|0\s*v.{0,20}\b0\b", blob)) or "0v" in blob.replace(" ", "")
    has_4095 = "4095" in blob
    has_mid = bool(re.search(r"1984|1985|2048|1[,.]6\s*v|1[,.]65", blob))
    has_bin = bool(re.search(r"binaire|0b[01]+|[01]{8,}", blob))
    q7 = 0
    if has_4095 or has_mid or re.search(r"n\s*=\s*u|4095\s*/\s*3", blob):
        q7 += 1
    if (has_4095 and has_mid) or (has_bin and has_4095):
        q7 += 1
    scores["TP2_Q7_Nbin"] = min(2, q7)

    # Q8 schéma Spice (TP2, pas la simu LED du TP1)
    tp2_spice = bool(
        re.search(
            r"(spice|simulateur).{0,80}(ldr|photor|gpio\s*22|gpio\s*34|analog|pont)|"
            r"(ldr|photor|gpio\s*22|analogread).{0,80}(spice|simulateur)|"
            r"tp2.{0,40}spice|spice.{0,40}tp2",
            blob,
        )
    )
    scores["TP2_Q8_Spice"] = 1 if tp2_spice else 0

    # Q9 tableau 3 situations (3 pts)
    volts = re.findall(r"\b\d+[.,]\d+\s*v\b|\b[0-3][.,]\d{1,3}\s*v\b", blob)
    situations = len(
        re.findall(
            r"lumi[eè]re forte|ombre|nuit|jour|faible|transition|limite|fen[eê]tre",
            blob,
        )
    )
    q9 = 0
    if situations >= 2 or (len(volts) >= 2 and re.search(r"ldr|photor|milieu|multim|tp2", blob)):
        q9 += 2
    if situations >= 3 or len(volts) >= 3 or (
        re.search(r"tableau", blob) and len(volts) >= 1 and re.search(r"ldr|lumi|nuit|jour", blob)
    ):
        q9 += 1
    scores["TP2_Q9_Tableau"] = min(3, q9)

    # Q10 vidéo tension / lumière (2) — fichier/lien vidéo + contexte TP2
    tp2_video_ctx = bool(
        re.search(
            r"tension|multim[eè]tre|ldr|photor|variation\s+tension|tp2|"
            r"gpio\s*22|photoresistance|/s/photoresistance|voltm",
            blob,
        )
    )
    if videos and tp2_video_ctx:
        scores["TP2_Q10_VideoU"] = 2
    elif videos and re.search(r"\btp2\b|photor|pont\s+diviseur|analogread", blob):
        scores["TP2_Q10_VideoU"] = 2
    elif re.search(r"vid[eé]o.{0,50}(tension|multim|ldr|photor|lumi[eè]re)", blob):
        scores["TP2_Q10_VideoU"] = 2  # lien/mention explicite → points
    else:
        scores["TP2_Q10_VideoU"] = 0

    # Q11 commenter le programme (2)
    if re.search(r"analogread", blob) and re.search(
        r"comment|explication|setup|loop|cette ligne|sert [aà]|permet", blob
    ):
        scores["TP2_Q11_Comment"] = 2
    elif "analogread" in blob and ("serial" in blob or "4095" in blob):
        scores["TP2_Q11_Comment"] = 1
    else:
        scores["TP2_Q11_Comment"] = 0

    # Q12 ADC/U pour 10000 / 10 / 1 lx (2)
    if re.search(r"10000\s*lx|10\s*000\s*lx|10\s*lx|1\s*lx|obscurit", blob) and re.search(
        r"adc|u22|tension|\d{2,4}", blob
    ):
        scores["TP2_Q12_SimLx"] = 2
    elif re.search(r"10000|10\s*lx|obscurit|lumi[eè]re vive", blob):
        scores["TP2_Q12_SimLx"] = 1
    else:
        scores["TP2_Q12_SimLx"] = 0

    # Q13 conclusion (1)
    if re.search(
        r"conclus|donc .{0,60}(lumi|adc|tension|augmente|diminue)|"
        r"(plus|moins)\s+(de\s+)?lumi.{0,40}(augmente|diminue|plus|moins)",
        blob,
    ):
        scores["TP2_Q13_Concl"] = 1
    else:
        scores["TP2_Q13_Concl"] = 0

    # Signal TP2 réel : au moins 2 critères distincts, ou total ≥ 3, ou « TP2 » explicite / ADC
    has_adc_tp2 = ("analogread" in blob) and re.search(
        r"\btp2\b|gpio\s*22|gpio\s*34|photor|ldr|4095|u22", blob
    )
    scored_keys = [k for k, v in scores.items() if v and not k.startswith("_")]
    tot = sum(scores[k] for k in scored_keys)
    # Schéma / spice / vidéo seuls (souvent présents dans un rendu TP1) ne valident pas le TP2
    theory_keys = [
        k
        for k in scored_keys
        if k not in ("TP2_Q1_Schema", "TP2_Q8_Spice", "TP2_Q10_VideoU")
    ]
    explicit = bool(re.search(r"\btp2\b|photor[eé]sistance|pont\s+diviseur", blob))
    strong = (
        bool(has_adc_tp2)
        or explicit
        or len(theory_keys) >= 2
        or tot >= 5
    )
    scores["_rendu"] = bool(strong)
    if not strong:
        for k in list(scores):
            if not k.startswith("_"):
                scores[k] = 0
    return scores


def match_student_key(nom: str, prenom: str, submission_keys: set[str]) -> str | None:
    k = norm_key(nom, prenom)
    if k in submission_keys:
        return k
    # try NOM_Prenom with first token only
    prenom0 = prenom.split()[0] if prenom else ""
    k2 = norm_key(nom, prenom0)
    if k2 in submission_keys:
        return k2
    # fuzzy: startswith nom_
    nom_u = re.sub(r"[^A-Z0-9\-]", "", nom.upper().replace("É", "E").replace("È", "E").replace("Ç", "C"))
    candidates = [s for s in submission_keys if s.startswith(nom_u + "_") or s.startswith(nom_u.replace("-", "") + "_")]
    if len(candidates) == 1:
        return candidates[0]
    # compound names
    for s in submission_keys:
        if nom_u.replace("-", "") in s.replace("-", "") and prenom0.upper().replace("É", "E")[:4] in s:
            return s
    return None


def style_header(cell, fill_color="1F4E79"):
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")


def build_sheet(wb, title: str, eleves: list, grades: dict):
    ws = wb.create_sheet(title)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    # Row 1: group headers
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws["A1"] = "Élève"
    style_header(ws["A1"])

    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    ws["B1"] = "Fichier déposé"
    style_header(ws["B1"], "375623")

    c = 3
    ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + len(TP1_COLS))
    ws.cell(1, c, f"TP1 — LED ESP32 (barème /{TP1_TOTAL} — sujet TP1_LED.pdf)")
    style_header(ws.cell(1, c), "C65911")
    for i, (key, mx, label) in enumerate(TP1_COLS):
        cell = ws.cell(2, c + i, f"{label}\n/{mx}")
        style_header(cell, "F4B183")
        cell.font = Font(bold=True, size=8, color="000000")
    tot1 = ws.cell(2, c + len(TP1_COLS), f"Total TP1\n/{TP1_TOTAL}")
    style_header(tot1, "C65911")

    c2 = 3 + len(TP1_COLS) + 1
    ws.merge_cells(start_row=1, start_column=c2, end_row=1, end_column=c2 + len(TP2_COLS))
    ws.cell(1, c2, f"TP2 — LDR & mesure (barème /{TP2_TOTAL} — sujet TP2_LDR_mesure.pdf)")
    style_header(ws.cell(1, c2), "2F5496")
    for i, (key, mx, label) in enumerate(TP2_COLS):
        cell = ws.cell(2, c2 + i, f"{label}\n/{mx}")
        style_header(cell, "8FAADC")
        cell.font = Font(bold=True, size=8, color="000000")
    tot2 = ws.cell(2, c2 + len(TP2_COLS), f"Total TP2\n/{TP2_TOTAL}")
    style_header(tot2, "2F5496")

    c3 = c2 + len(TP2_COLS) + 1
    grand = TP1_TOTAL + TP2_TOTAL
    ws.merge_cells(start_row=1, start_column=c3, end_row=2, end_column=c3)
    ws.cell(1, c3, f"Total\n/{grand}")
    style_header(ws.cell(1, c3), "833C0C")

    c4 = c3 + 1
    ws.merge_cells(start_row=1, start_column=c4, end_row=2, end_column=c4)
    ws.cell(1, c4, "Commentaire / correction")
    style_header(ws.cell(1, c4), "595959")

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 55

    eleves_sorted = sorted(eleves, key=lambda e: (e.get("nom", ""), e.get("prenom", "")))
    for r, e in enumerate(eleves_sorted, start=3):
        nom = f"{e.get('nom', '')} {e.get('prenom', '')}".strip()
        g = grades.get(e["_match_key"] or "", {})
        ws.cell(r, 1, nom).border = thin
        fname = g.get("filename", "— non rendu —")
        ws.cell(r, 2, fname).border = thin

        col = 3
        s1 = 0
        for key, mx, _ in TP1_COLS:
            val = g.get(key)
            cell = ws.cell(r, col, val if g.get("tp1_rendu") else "")
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
            if isinstance(val, (int, float)):
                s1 += val
            col += 1
        cell = ws.cell(r, col, s1 if g.get("tp1_rendu") else "")
        cell.border = thin
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        col += 1

        s2 = 0
        for key, mx, _ in TP2_COLS:
            val = g.get(key)
            cell = ws.cell(r, col, val if g.get("tp2_rendu") else "")
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
            if isinstance(val, (int, float)) and g.get("tp2_rendu"):
                s2 += val
            col += 1
        cell = ws.cell(r, col, s2 if g.get("tp2_rendu") else "")
        cell.border = thin
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        col += 1

        total = (s1 if g.get("tp1_rendu") else 0) + (s2 if g.get("tp2_rendu") else 0)
        cell = ws.cell(r, col, total if (g.get("tp1_rendu") or g.get("tp2_rendu")) else "")
        cell.border = thin
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        col += 1

        ws.cell(r, col, g.get("comment", "")).border = thin

        if not g.get("tp1_rendu") and not g.get("tp2_rendu"):
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="FCE4D6")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36
    for i in range(3, c4 + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if i < c4 else 40
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(c4)}{2 + len(eleves_sorted)}"


def dedupe_subs(items: list) -> list:
    """Garde un exemplaire unique par (name, size) — local prioritaire sur backup."""
    seen = {}
    for it in items:
        k = (it["name"], it["size"])
        if k not in seen or it["kind"] == "local":
            seen[k] = it
    return list(seen.values())


def main():
    eleves = load_eleves()
    latest, all_subs = collect_submissions()
    print(f"Élèves avec dépôt: {len(all_subs)}")

    # Grade ALL submissions per student (TP1 + TP2 often in separate uploads)
    graded_by_key = {}
    for key, items in sorted(all_subs.items()):
        items = dedupe_subs(items)
        contents = []
        names = []
        for sub in sorted(items, key=lambda x: x["date"]):
            print(f"Analyse {sub['name']} ({sub['kind']}) ...")
            try:
                content = extract_content(sub)
            except Exception as ex:
                content = {
                    "filename": sub["name"],
                    "text": "",
                    "images": 0,
                    "videos": 0,
                    "inos": [],
                    "size": sub["size"],
                }
                print("  ERREUR lecture:", ex)
            contents.append(content)
            names.append(sub["name"])
        merged = merge_contents(contents)
        s1 = score_tp1(merged)
        s2 = score_tp2(merged)
        comment_parts = []
        if not s1.get("_rendu"):
            comment_parts.append("TP1 non identifiable / non rendu")
        else:
            if s1.get("TP1_Q5_PhotoMontage", 0) < 2:
                comment_parts.append("TP1: photo montage manquante ou douteuse")
            if s1.get("TP1_Q6_VideoCligno", 0) < 2:
                comment_parts.append("TP1: vidéo clignotement manquante")
            if s1.get("TP1_Q4_ProgBlink", 0) < 3:
                comment_parts.append("TP1: prog. clignotement 500 ms incomplet")
        if not s2.get("_rendu"):
            comment_parts.append("TP2 non identifiable / non rendu")
        else:
            if s2.get("TP2_Q9_Tableau", 0) < 2:
                comment_parts.append("TP2: tableau mesures incomplet")
            if s2.get("TP2_Q10_VideoU", 0) < 1:
                comment_parts.append("TP2: vidéo tension/lumière manquante")
        graded_by_key[key] = {
            "filename": " | ".join(names),
            "tp1_rendu": s1.get("_rendu", False),
            "tp2_rendu": s2.get("_rendu", False),
            "comment": " ; ".join(comment_parts),
            **{k: v for k, v in s1.items() if not k.startswith("_")},
            **{k: v for k, v in s2.items() if not k.startswith("_")},
        }
        print(
            f"  -> TP1={'oui' if s1.get('_rendu') else 'non'} "
            f"sum={sum(v for k,v in s1.items() if not k.startswith('_'))}/{TP1_TOTAL} | "
            f"TP2={'oui' if s2.get('_rendu') else 'non'} "
            f"sum={sum(v for k,v in s2.items() if not k.startswith('_'))}/{TP2_TOTAL}"
        )

    # Match students
    sub_keys = set(latest.keys())
    for classe, liste in eleves.items():
        for e in liste:
            mk = match_student_key(e.get("nom", ""), e.get("prenom", ""), sub_keys)
            e["_match_key"] = mk
            if not mk:
                print(f"Sans copie: {classe} {e.get('nom')} {e.get('prenom')}")

    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    # Bareme sheet
    ws = wb.create_sheet("Barème", 0)
    ws["A1"] = "Barème issu des sujets TP1_LED.pdf et TP2_LDR_mesure.pdf"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A3"] = "TP1 — Allumer une LED avec l'ESP32 (TP1_LED.pdf)"
    ws["A3"].font = Font(bold=True, color="C65911")
    row = 4
    for key, mx, label in TP1_COLS:
        ws.cell(row, 1, label)
        ws.cell(row, 2, f"/{mx}")
        row += 1
    ws.cell(row, 1, "Total TP1")
    ws.cell(row, 2, f"/{TP1_TOTAL}")
    ws.cell(row, 1).font = Font(bold=True)
    row += 2
    ws.cell(row, 1, "TP2 — Mesurer la luminosité avec une LDR (TP2_LDR_mesure.pdf)")
    ws.cell(row, 1).font = Font(bold=True, color="2F5496")
    row += 1
    for key, mx, label in TP2_COLS:
        ws.cell(row, 1, label)
        ws.cell(row, 2, f"/{mx}")
        row += 1
    ws.cell(row, 1, "Total TP2")
    ws.cell(row, 2, f"/{TP2_TOTAL}")
    ws.cell(row, 1).font = Font(bold=True)
    row += 2
    ws.cell(row, 1, "Total général")
    ws.cell(row, 2, f"/{TP1_TOTAL + TP2_TOTAL}")
    ws.cell(row, 1).font = Font(bold=True)
    row += 2
    ws.cell(
        row,
        1,
        "Correction assistée d'après les dépôts (texte, code, photos, vidéos). "
        "Points photo/vidéo accordés si fichier image/vidéo présent ou lien explicite. "
        "À relire / ajuster si besoin.",
    )
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 10

    build_sheet(wb, "1ère STI2D 1", eleves.get("1ere_STI2D1", []), graded_by_key)
    build_sheet(wb, "1ère STI2D 2", eleves.get("1ere_STI2D2", []), graded_by_key)

    # Also map grades for unmatched via student match keys used in sheet
    # Fix: build_sheet looks up grades.get(e["_match_key"]) — need to pass keyed by match key
    # graded_by_key already keyed by submission key; _match_key points to that. OK.

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(OUT_XLSX)
        print("Excel écrit:", OUT_XLSX)
    except PermissionError:
        alt = OUT_XLSX.with_name(OUT_XLSX.stem + "_corrige.xlsx")
        wb.save(alt)
        print("Fichier verrouillé (ouvert dans Excel). Écrit:", alt)


if __name__ == "__main__":
    main()
