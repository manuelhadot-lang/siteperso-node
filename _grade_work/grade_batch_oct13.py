# -*- coding: utf-8 -*-
"""Corrige les rendus du 13/10 et met à jour Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx."""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import build_notes_excel as b
from openpyxl import load_workbook
from openpyxl.styles import Font

DL = Path(r"c:\Users\mhado\Downloads")
ATT = Path(
    r"c:\Users\mhado\.cursor\projects\c-Users-mhado-Desktop-site-site-site-site-site-sitepersoLight-1-sitepersoLight-siteperso-siteperso-siteperso-siteperso-node-main-siteperso-node-main\attachments\1269adb1-bd26-4d68-a15b-caabcf063aa7"
)
ORIG = b.OUT_XLSX
MAJ = ORIG.with_name("Notes_TP1_TP2_1ere_STI2D_maj_TP2.xlsx")

# Chaque entrée: key, also[], paths (list Path) — plusieurs fichiers fusionnés (odt+mov…)
BATCH = [
    {"key": "BAZIN_FRANCOIS", "also": ["CHEVANNE-CARENCO_GABRIEL"], "paths": [DL / "BAZIN_Francois_2026-10-13.zip"]},
    # CHEVANNE même zip contenu — noté via also de BAZIN; on garde aussi entrée dédiée si besoin
    {"key": "BODILIS_PAUL", "also": ["GIRCOURT_MATHIEU"], "paths": [DL / "BODILIS_Paul_2026-10-13.zip"]},
    {"key": "COJANU_ALEXANDRE", "also": ["MEUDIC_THOMAS"], "paths": [DL / "COJANU_Alexandre_2026-10-13 (1).zip"]},
    {
        "key": "DHOKERS-ALLUSSON_ADAM",
        "also": [],
        "match_names": ["D'HOKERS", "DHOKERS", "HOKERS"],
        "paths": [
            DL / "D'HOKERS-ALLUSSON_Adam_2026-10-13.odt",
            DL / "D'HOKERS-ALLUSSON_Adam_2026-10-13 (1).mov",
        ],
    },
    {
        "key": "FEAUGAS_JEAN_CHARLES",
        "also": [],
        "paths": [DL / "FEAUGAS_Jean_Charles_2026-10-13 (1).odt"],
    },
    {"key": "LEMONDE-SAN_SORYA", "also": [], "paths": [DL / "LEMONDE-SAN_Sorya_2026-10-13.zip"]},
    {"key": "RACHI_MARWANE", "also": [], "paths": [DL / "RACHI_Marwane_2026-10-13.docx"]},
    {"key": "VARILLON_AMAURY", "also": [], "paths": [DL / "VARILLON_Amaury_2026-10-13.zip"]},
    {"key": "LEYE_MARIAMA", "also": [], "paths": [DL / "LEYE_Mariama_2026-10-13.pdf"]},
    {"key": "RECHE_VICTOR", "also": [], "paths": [DL / "RECHE_Victor_2026-10-13.pdf"]},
    {
        "key": "RIGAMONTI_CLEMENT",
        "also": [],
        "paths": list(DL.glob("RIGAMONTI_Cl*_2026-10-13.pdf"))
        or list(ATT.glob("RIGAMONTI_Cl*_2026-10-13.pdf")),
    },
]


def content_from_path(p: Path) -> dict:
    if not p.exists():
        raise FileNotFoundError(p)
    if p.is_dir():
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in p.rglob("*"):
                if f.is_file():
                    zf.write(f, arcname=str(f.relative_to(p)).replace("\\", "/"))
        zinfo = b.analyze_zip_bytes(buf.getvalue())
        return {
            "filename": p.name + "/",
            "size": sum(f.stat().st_size for f in p.rglob("*") if f.is_file()),
            "text": p.name + "\n" + zinfo["text"],
            "images": zinfo["images"],
            "videos": zinfo.get("videos", 0),
            "inos": zinfo.get("inos", []),
            "ext": "dir",
        }
    name = p.name.lower()
    data = p.read_bytes()
    out = {
        "filename": p.name,
        "size": p.stat().st_size,
        "text": p.name + "\n",
        "images": 0,
        "videos": 0,
        "inos": [],
        "ext": p.suffix.lower(),
    }
    if name.endswith(".docx"):
        d = b.analyze_docx_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".odt"):
        d = b.analyze_odt_bytes(data)
        out["text"] += d["text"]
        out["images"] = d.get("images", 0)
        out["videos"] = d.get("videos", 0)
    elif name.endswith(".pdf"):
        out["text"] += b.read_pdf_text(data)
    elif name.endswith(".zip"):
        zinfo = b.analyze_zip_bytes(data)
        out["text"] += zinfo["text"]
        out["images"] = zinfo["images"]
        out["videos"] = zinfo.get("videos", 0)
        out["inos"] = zinfo.get("inos", [])
    elif name.endswith((".mp4", ".mov", ".m4v", ".avi")):
        out["videos"] = 1
        out["text"] += f"\n[video-file:{p.name}]"
    else:
        out["text"] += data.decode("utf-8", errors="ignore")
    return out


def looks_like_tp2_only(content: dict) -> bool:
    blob = (content.get("text") or "").lower()
    name = (content.get("filename") or "").lower()
    if re.search(r"\btp\s*2\b|/tp 2/|tp2[_ .]", name):
        return True
    if "tp2" in blob and "tp1" not in blob[:500]:
        # faible
        pass
    # dossier/archive clairement TP2
    if re.search(r"(^|\n)tp\s*2/|tp2\.odt|tp2 final|tp2_", blob):
        return True
    return False


def norm_cell(s: str) -> str:
    s = (s or "").upper()
    for a, c in [
        ("É", "E"), ("È", "E"), ("Ê", "E"), ("Ë", "E"),
        ("À", "A"), ("Â", "A"), ("Ù", "U"), ("Û", "U"),
        ("Ô", "O"), ("Î", "I"), ("Ï", "I"), ("Ç", "C"), ("Œ", "OE"),
        ("'", ""), ("’", ""), ("-", ""), (" ", ""),
    ]:
        s = s.replace(a, c)
    return s


def find_row(ws, key: str, match_names: list[str] | None = None):
    key_n = norm_cell(key.replace("_", " "))
    for r in range(3, ws.max_row + 1):
        cell = norm_cell(str(ws.cell(r, 1).value or ""))
        if not cell:
            continue
        if match_names:
            if any(norm_cell(m) in cell for m in match_names):
                return r
        # NOM + début prénom
        parts = key.split("_")
        nom = norm_cell(parts[0])
        prenom = norm_cell(parts[-1]) if len(parts) > 1 else ""
        if nom and nom in cell and (not prenom or prenom[:4] in cell):
            return r
    return None


def apply_scores(ws, row: int, s1: dict, s2: dict, label: str, prefer_keep_tp1: bool):
    tp1_n = len(b.TP1_COLS)
    tp2_n = len(b.TP2_COLS)
    col_tp1 = 3
    col_tp1_tot = 3 + tp1_n
    col_tp2 = col_tp1_tot + 1
    col_tp2_tot = col_tp2 + tp2_n
    col_grand = col_tp2_tot + 1
    col_comment = col_grand + 1

    prev_f = ws.cell(row, 2).value or ""
    if label not in str(prev_f):
        ws.cell(row, 2, str(prev_f) + " + " + label)

    # TP1
    write_tp1 = bool(s1.get("_rendu"))
    if write_tp1 and prefer_keep_tp1:
        try:
            old = float(ws.cell(row, col_tp1_tot).value or 0)
        except Exception:
            old = 0
        new = sum(v for k, v in s1.items() if not k.startswith("_"))
        if old > new:
            write_tp1 = False
            print(f"    conserve TP1 {old} > {new}")

    if write_tp1:
        t1 = 0
        for i, (k, mx, _) in enumerate(b.TP1_COLS):
            val = s1.get(k, 0)
            ws.cell(row, col_tp1 + i, val)
            t1 += val
        ws.cell(row, col_tp1_tot, t1)
        ws.cell(row, col_tp1_tot).font = Font(bold=True)

    if s2.get("_rendu"):
        t2 = 0
        for i, (k, mx, _) in enumerate(b.TP2_COLS):
            val = s2.get(k, 0)
            ws.cell(row, col_tp2 + i, val)
            t2 += val
        ws.cell(row, col_tp2_tot, t2)
        ws.cell(row, col_tp2_tot).font = Font(bold=True)

    try:
        t1f = float(ws.cell(row, col_tp1_tot).value or 0)
    except Exception:
        t1f = 0
    try:
        t2f = float(ws.cell(row, col_tp2_tot).value or 0)
    except Exception:
        t2f = 0
    if ws.cell(row, col_tp1_tot).value not in (None, "") or ws.cell(row, col_tp2_tot).value not in (None, ""):
        ws.cell(row, col_grand, t1f + t2f)
        ws.cell(row, col_grand).font = Font(bold=True)

    bits = []
    if write_tp1:
        bits.append(f"TP1={sum(v for k,v in s1.items() if not k.startswith('_'))}/20")
    if s2.get("_rendu"):
        bits.append(f"TP2={sum(v for k,v in s2.items() if not k.startswith('_'))}/20")
    note = f"Corrigé {label}: " + " ; ".join(bits) if bits else f"Vu {label}"
    com = str(ws.cell(row, col_comment).value or "")
    if s1.get("_rendu") and write_tp1:
        com = com.replace("TP1 non identifiable / non rendu", "").strip(" ;")
    if s2.get("_rendu"):
        com = com.replace("TP2 non identifiable / non rendu", "").strip(" ;")
    if note not in com:
        ws.cell(row, col_comment, (com + " ; " if com else "") + note)


def main():
    # Aussi noter CHEVANNE / GIRCOURT / MEUDIC via also — pas besoin d'analyser 2x les gros zip
    results = []
    for item in BATCH:
        paths = [p for p in item["paths"] if p and Path(p).exists()]
        print(f"\n=== {item['key']} paths={len(paths)} ===")
        if not paths:
            print("  MANQUANT", item["paths"])
            continue
        contents = [content_from_path(Path(p)) for p in paths]
        merged = b.merge_contents(contents)
        s1 = b.score_tp1(merged)
        s2 = b.score_tp2(merged)
        tp2_only = looks_like_tp2_only(merged)
        if tp2_only and s2.get("_rendu"):
            # annuler faux points TP1 (vidéos TP2)
            for k in list(s1):
                if not k.startswith("_"):
                    s1[k] = 0
            s1["_rendu"] = False
            print("  (detecte TP2-only -> ignore score TP1)")
        t1 = sum(v for k, v in s1.items() if not k.startswith("_"))
        t2 = sum(v for k, v in s2.items() if not k.startswith("_"))
        print(f"  files: {[c['filename'] for c in contents]}")
        print(f"  imgs={merged['images']} vids={merged['videos']}")
        print(f"  TP1={'oui' if s1.get('_rendu') else 'non'} {t1}/20 | TP2={'oui' if s2.get('_rendu') else 'non'} {t2}/20")
        print("  TP2 nz", {k: v for k, v in s2.items() if not k.startswith("_") and v})
        results.append({**item, "s1": s1, "s2": s2, "t1": t1, "t2": t2, "label": " + ".join(c["filename"] for c in contents), "tp2_only": tp2_only})

    if not MAJ.exists():
        raise SystemExit(f"Maj Excel manquant: {MAJ}")
    wb = load_workbook(MAJ)

    for res in results:
        keys = [res["key"]] + list(res.get("also") or [])
        for sheet in wb.sheetnames:
            if sheet.startswith("Barème") or sheet.startswith("TP2") or sheet.startswith("Ajouts"):
                continue
            ws = wb[sheet]
            for key in keys:
                row = find_row(ws, key, res.get("match_names"))
                if not row:
                    continue
                prefer = not res["tp2_only"]  # si TP2 only, ne pas écrire TP1
                # also-keys: same scores
                apply_scores(ws, row, res["s1"], res["s2"], res["label"], prefer_keep_tp1=True)
                print(f"  OK {sheet} r{row} {ws.cell(row,1).value} <- {key} TP2={res['t2']}")

    name = "Ajouts 13 oct"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, 0)
    ws["A1"] = "Rendus 13/10 (original Notes_TP1_TP2_1ere_STI2D.xlsx non modifié)"
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Élèves", "Fichier(s)", "TP1 /20", "TP2 /20"])
    for res in results:
        eleves = ", ".join([res["key"]] + list(res.get("also") or []))
        ws.append([
            eleves,
            res["label"],
            res["t1"] if res["s1"].get("_rendu") else "— (conservé / N/A)",
            res["t2"] if res["s2"].get("_rendu") else "—",
        ])
    ws.append([])
    ws.append(["Non récupérés (site protégé par mot de passe)", "ROUVEL Gabin + VERDYS Maxime 13/10.pdf"])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12

    wb.save(MAJ)
    print("\nSauvé:", MAJ)
    print("Original intact:", ORIG)


if __name__ == "__main__":
    main()
